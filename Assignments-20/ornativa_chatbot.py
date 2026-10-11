from __future__ import annotations

import csv
import difflib
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from dotenv import load_dotenv

load_dotenv(Path(__file__).resolve().parent / ".env")
load_dotenv(Path(__file__).resolve().parent.parent / ".env")

try:
    from openai import OpenAI
except Exception:
    OpenAI = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    import pandas as pd
except Exception:
    pd = None

try:
    from langchain_community.vectorstores import FAISS
    from langchain_openai import OpenAIEmbeddings
except Exception:
    FAISS = None
    OpenAIEmbeddings = None


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------

@dataclass
class Record:
    source: str
    record_id: str
    fields: Dict[str, Any]

    def get(self, key: str, default: Any = None) -> Any:
        if key in self.fields:
            return self.fields[key]
        wanted = normalise_key(key)
        for actual, value in self.fields.items():
            if normalise_key(actual) == wanted:
                return value
        return default

    def searchable_text(self) -> str:
        return " | ".join(f"{k}: {v}" for k, v in self.fields.items())


# ---------------------------------------------------------------------------
# Normalisation helpers
# ---------------------------------------------------------------------------

KEY_ALIASES = {
    "productid": "product_id",
    "product_id": "product_id",
    "id": "product_id",
    "sku": "product_id",
    "name": "name",
    "productname": "name",
    "product_name": "name",
    "title": "name",
    "category": "category",
    "type": "category",
    "material": "material",
    "metal": "material",
    "weight": "weight",
    "price": "price",
    "cost": "price",
    "amount": "price",
    "stock": "stock",
    "availability": "stock",
    "status": "stock",
}


def normalise_key(value: Any) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return KEY_ALIASES.get(text, text)


def normalise_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return str(value).strip() if not isinstance(value, (dict, list)) else value


def parse_number(value: Any) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).lower().replace(",", "").replace("₹", "").replace("rs.", "").replace("inr", "")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    number = float(match.group())
    if "crore" in text or "crores" in text:
        number *= 10_000_000
    elif "lakh" in text or "lakhs" in text or "lac" in text or "lacs" in text:
        number *= 100_000
    elif "thousand" in text or "thousands" in text:
        number *= 1_000
    return number


def canonical_material(text: str) -> str:
    text = str(text or "").lower()
    match = re.search(r"\b(18|20|22|24)\s*k(?:\s*gold)?\b", text)
    if match:
        return f"{match.group(1)}K Gold"
    return str(text).strip()


def normalise_record(source: str, index: int, fields: Dict[str, Any]) -> Record:
    cleaned = {}
    for key, value in fields.items():
        canonical = normalise_key(key)
        if canonical:
            cleaned[canonical] = normalise_value(value)

    if "material" in cleaned:
        cleaned["material"] = canonical_material(cleaned["material"])

    record_id = str(cleaned.get("product_id") or cleaned.get("name") or f"{Path(source).name}:{index}")
    return Record(source=source, record_id=record_id, fields=cleaned)


# ---------------------------------------------------------------------------
# Ingestion: PDF / CSV / JSON / XLSX / DOCX / TXT
# ---------------------------------------------------------------------------

class UniversalIngestor:
    """
    Converts common document formats into a common list[Record].

    Structured sources are parsed into fields. Unstructured text remains
    available for semantic retrieval through the text_records collection.
    """

    def __init__(self, source_path: str | Path):
        self.source_path = Path(source_path)

    def load(self) -> Tuple[List[Record], List[Record]]:
        files = self._files()
        structured: List[Record] = []
        text_records: List[Record] = []

        for path in files:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                s, t = self._pdf(path)
            elif suffix == ".csv":
                s, t = self._csv(path)
            elif suffix == ".json":
                s, t = self._json(path)
            elif suffix in {".xlsx", ".xlsm"}:
                s, t = self._xlsx(path)
            elif suffix == ".docx":
                s, t = self._docx(path)
            elif suffix in {".txt", ".md"}:
                s, t = self._text(path)
            else:
                continue
            structured.extend(s)
            text_records.extend(t)

        if not structured and not text_records:
            raise ValueError(f"No supported/readable documents found in {self.source_path}")

        return structured, text_records

    def _files(self) -> List[Path]:
        if self.source_path.is_file():
            return [self.source_path]
        if self.source_path.is_dir():
            supported = {".pdf", ".csv", ".json", ".xlsx", ".xlsm", ".docx", ".txt", ".md"}
            return sorted(p for p in self.source_path.iterdir() if p.suffix.lower() in supported)
        raise FileNotFoundError(self.source_path)

    def _pdf(self, path: Path):
        if PdfReader is None:
            raise ImportError("Install pypdf to read PDF files.")
        raw = []
        for page_no, page in enumerate(PdfReader(str(path)).pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                raw.append((page_no, text))

        full = "\n".join(text for _, text in raw)
        records = self._parse_product_blocks(path.name, full)

        texts = [
            Record(
                source=path.name,
                record_id=f"{path.name}:page:{page_no}",
                fields={"text": text, "page": page_no},
            )
            for page_no, text in raw
        ]
        return records, texts

    def _parse_product_blocks(self, source: str, text: str) -> List[Record]:
        # Supports the common "ProductID: ... Name: ... Category: ..."
        # catalogue shape without making the whole engine dependent on it.
        clean = re.sub(r"\s+", " ", text).strip()
        pattern = re.compile(
            r"ProductID:\s*(?P<product_id>[A-Za-z0-9_-]+)\s+"
            r"Name:\s*(?P<name>.+?)\s+"
            r"Category:\s*(?P<category>.+?)\s+"
            r"Material:\s*(?P<material>.+?)\s+"
            r"Weight:\s*(?P<weight>.+?)\s+"
            r"Price:\s*(?P<price>[\d,.]+)\s+"
            r"Stock:\s*(?P<stock>Available|Out of Stock)",
            re.IGNORECASE,
        )
        records = []
        for i, match in enumerate(pattern.finditer(clean), 1):
            data = match.groupdict()
            data["price"] = parse_number(data["price"])
            records.append(normalise_record(source, i, data))
        return records

    def _csv(self, path: Path):
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))
        records = [
            normalise_record(path.name, i, row)
            for i, row in enumerate(rows, 1)
        ]
        return records, []

    def _json(self, path: Path):
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            if isinstance(data.get("products"), list):
                data = data["products"]
            else:
                data = [data]
        if not isinstance(data, list):
            return [], [Record(path.name, f"{path.name}:1", {"text": json.dumps(data)})]
        records = [
            normalise_record(path.name, i, item)
            for i, item in enumerate(data, 1)
            if isinstance(item, dict)
        ]
        return records, []

    def _xlsx(self, path: Path):
        if load_workbook is None:
            raise ImportError("Install openpyxl to read Excel files.")
        wb = load_workbook(path, read_only=True, data_only=True)
        records = []
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                headers = [str(x or "").strip() for x in next(rows)]
            except StopIteration:
                continue
            for i, values in enumerate(rows, 2):
                row = {
                    headers[j]: values[j]
                    for j in range(min(len(headers), len(values)))
                    if headers[j]
                }
                if any(v not in (None, "") for v in row.values()):
                    records.append(normalise_record(f"{path.name}:{ws.title}", i, row))
        return records, []

    def _docx(self, path: Path):
        if DocxDocument is None:
            raise ImportError("Install python-docx to read DOCX files.")
        doc = DocxDocument(str(path))
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables = []
        for table in doc.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if rows:
                tables.append(rows)

        records = []
        for table_index, rows in enumerate(tables, 1):
            headers = rows[0]
            for row_index, values in enumerate(rows[1:], 2):
                data = {
                    headers[i]: values[i]
                    for i in range(min(len(headers), len(values)))
                    if headers[i]
                }
                records.append(normalise_record(f"{path.name}:table:{table_index}", row_index, data))

        text = "\n".join(lines)
        text_records = []
        if text:
            text_records.append(Record(path.name, f"{path.name}:text", {"text": text}))
        return records, text_records

    def _text(self, path: Path):
        text = path.read_text(encoding="utf-8", errors="replace").strip()
        return [], [Record(path.name, f"{path.name}:text", {"text": text})] if text else []


# ---------------------------------------------------------------------------
# Generic query planning
# ---------------------------------------------------------------------------

QUERY_SYSTEM = """
You are a query planner for a document-grounded assistant.

Convert the user's question into a small, generic JSON query plan.
Do not answer the question.

Allowed operation values:
- "filter"
- "lookup"
- "compare"
- "aggregate"
- "semantic"
- "answer_from_context"

Allowed fields are discovered from the catalogue and may include:
product_id, name, category, material, weight, price, stock, plus other fields.

Known field values are supplied with the request. When a categorical value is
needed, use the exact supplied value rather than inventing or paraphrasing one.

For filter conditions use:
{"field": "...", "operator": "eq|neq|contains|gt|gte|lt|lte", "value": ...}

For sorting:
{"field": "...", "direction": "asc|desc"}

For output fields, use a list of canonical field names.

Output rules:
- For product selection, filtering, and ranking requests, return ["name"] only.
- When the user explicitly asks for a field, return ["name", "<requested field>"] only.
- When the user asks for "details" without naming fields, return
    {"clarify_fields": true} and no output fields.

For aggregate use:
{"function": "count|min|max|sum|avg", "field": "..."}

If the user says "all", "every", "which products", "show products", or asks
for a constrained list, use operation "filter" and do not impose a result limit.

For follow-ups like "what are their prices?", preserve the prior selected
record IDs if supplied in the conversation context.

Return JSON only.
"""


class QueryPlanner:
    def __init__(self):
        self.client = OpenAI(api_key=os.getenv("OPENAI_API_KEY")) if OpenAI and os.getenv("OPENAI_API_KEY") else None

    def plan(
        self,
        question: str,
        fields: Iterable[str] | Dict[str, List[Any]],
        previous_ids: Optional[List[str]] = None,
    ) -> Dict[str, Any]:
        field_values = fields if isinstance(fields, dict) else {}
        field_list = sorted(set(field_values or fields))
        if self.client:
            try:
                prompt = (
                    QUERY_SYSTEM
                    + "\nAvailable fields:\n"
                    + json.dumps(field_list)
                    + "\nKnown field values:\n"
                    + json.dumps(field_values)
                    + "\nPrevious selected record IDs:\n"
                    + json.dumps(previous_ids or [])
                    + "\nUser question:\n"
                    + question
                )
                response = self.client.responses.create(
                    model=os.getenv("OPENAI_QUERY_MODEL", "gpt-4o-mini"),
                    input=[{"role": "user", "content": prompt}],
                )
                raw = response.output_text.strip()
                model_plan = self._validate_plan(json.loads(raw), field_list, previous_ids or [])
                fallback_plan = self._fallback(question, field_list, previous_ids or [], field_values)
                return self._merge_explicit_constraints(model_plan, fallback_plan)
            except Exception:
                pass

        return self._fallback(question, field_list, previous_ids or [], field_values)

    @staticmethod
    def _merge_explicit_constraints(
        model_plan: Dict[str, Any],
        fallback_plan: Dict[str, Any],
    ) -> Dict[str, Any]:
        """Retain filters that are explicit in the user's wording.

        The LLM may validly return JSON while silently dropping one of several
        catalogue constraints. The deterministic parser recognizes direct
        material, category, price, and stock terms and restores only missing
        field constraints before the engine executes the plan.
        """
        merged = dict(model_plan)
        conditions = list(model_plan.get("conditions", []))
        constrained_fields = {condition.get("field") for condition in conditions}
        for condition in fallback_plan.get("conditions", []):
            if condition.get("field") not in constrained_fields:
                conditions.append(condition)
                constrained_fields.add(condition.get("field"))
        merged["conditions"] = conditions

        if fallback_plan.get("clarify_fields"):
            merged["clarify_fields"] = True
            merged["output_fields"] = []
        elif fallback_plan.get("sort") and not merged.get("sort"):
            merged["sort"] = fallback_plan["sort"]
            merged["limit"] = fallback_plan.get("limit")
            merged["output_fields"] = ["name"]
        elif fallback_plan.get("conditions") and fallback_plan.get("output_fields") == ["name"]:
            merged["output_fields"] = ["name"]

        return merged

    @staticmethod
    def _validate_plan(
        plan: Any,
        fields: List[str],
        previous_ids: List[str],
    ) -> Dict[str, Any]:
        """Keep model output within the query engine's supported contract."""
        if not isinstance(plan, dict):
            raise ValueError("Query plan must be a JSON object.")

        valid_operations = {"filter", "lookup", "compare", "aggregate", "semantic", "answer_from_context"}
        valid_operators = {"eq", "neq", "contains", "gt", "gte", "lt", "lte"}
        valid_fields = set(fields)
        operation = plan.get("operation", "filter")
        if operation not in valid_operations:
            raise ValueError("Unsupported query operation.")

        conditions = []
        for condition in plan.get("conditions", []):
            if not isinstance(condition, dict):
                continue
            field = normalise_key(condition.get("field"))
            operator = condition.get("operator")
            if field in valid_fields and operator in valid_operators and "value" in condition:
                conditions.append({"field": field, "operator": operator, "value": condition["value"]})

        output_fields = [
            normalise_key(field)
            for field in plan.get("output_fields", [])
            if normalise_key(field) in valid_fields
        ]
        sort = plan.get("sort")
        if not isinstance(sort, dict) or normalise_key(sort.get("field")) not in valid_fields:
            sort = None
        else:
            direction = sort.get("direction")
            sort = {"field": normalise_key(sort["field"]), "direction": direction if direction in {"asc", "desc"} else "asc"}

        validated = {
            "operation": operation,
            "conditions": conditions,
            "sort": sort,
            "output_fields": output_fields or ["name"],
        }
        if plan.get("clarify_fields"):
            validated["clarify_fields"] = True
            validated["output_fields"] = []
        if isinstance(plan.get("limit"), int) and plan["limit"] >= 0:
            validated["limit"] = plan["limit"]
        if plan.get("record_ids"):
            validated["record_ids"] = [str(record_id) for record_id in plan["record_ids"] if str(record_id) in previous_ids]
        if isinstance(plan.get("aggregate"), dict):
            aggregate = plan["aggregate"]
            if aggregate.get("function") in {"count", "min", "max", "sum", "avg"} and normalise_key(aggregate.get("field")) in valid_fields:
                validated["aggregate"] = {"function": aggregate["function"], "field": normalise_key(aggregate["field"])}
        return validated

    def _fallback(
        self,
        question: str,
        fields: List[str],
        previous_ids: List[str],
        field_values: Dict[str, List[Any]] | None = None,
    ) -> Dict[str, Any]:
        q = question.lower()
        plan: Dict[str, Any] = {
            "operation": "filter",
            "conditions": [],
            "sort": None,
            "output_fields": [],
        }

        if previous_ids and any(x in q for x in ["these", "their", "them", "those", "previous", "above"]):
            plan["record_ids"] = previous_ids

        for field, values in (field_values or {}).items():
            if field in {"price", "weight"}:
                continue
            matched_value = self._resolve_catalogue_value(q, values)
            if matched_value is not None:
                plan["conditions"].append({
                    "field": field,
                    "operator": "eq",
                    "value": matched_value,
                })

        price = re.search(
            r"\b(under|below|less than|up to|at most|above|over|more than|at least)\s*"
            r"(?:₹|rs\.?|inr|rupees?)?\s*([\d,.]+)\s*(lakh|lakhs|thousand|thousands|crore|crores)?",
            q,
        )
        if price:
            amount = parse_number(price.group(2) + " " + (price.group(3) or ""))
            op = price.group(1)
            operator = "lte" if op in {"under", "below", "less than", "up to", "at most"} else "gte"
            plan["conditions"].append({"field": "price", "operator": operator, "value": amount})

        if "cheapest" in q or "lowest price" in q:
            plan["sort"] = {"field": "price", "direction": "asc"}
            plan["limit"] = 1
        elif "most expensive" in q or "highest price" in q:
            plan["sort"] = {"field": "price", "direction": "desc"}
            plan["limit"] = 1
        elif "lightest" in q or "lowest weight" in q or "least weight" in q:
            plan["sort"] = {"field": "weight", "direction": "asc"}
            plan["limit"] = 1
        elif "heaviest" in q or "highest weight" in q:
            plan["sort"] = {"field": "weight", "direction": "desc"}
            plan["limit"] = 1

        if "detail" in q and not any(field in q for field in ["price", "cost", "weight", "stock", "availability", "material", "category"]):
            plan["clarify_fields"] = True
        elif plan["sort"] is None:
            requested_fields = [
                field for field in fields
                if re.search(rf"\b{re.escape(field.replace('_', ' '))}\b", q)
                and not any(condition.get("field") == field for condition in plan["conditions"])
            ]
            if requested_fields:
                plan["output_fields"] = ["name", *requested_fields]
            else:
                plan["output_fields"] = ["name"]
        else:
            plan["output_fields"] = ["name"]

        return plan

    @staticmethod
    def _resolve_catalogue_value(question: str, values: List[Any]) -> Any | None:
        """Resolve one unambiguous catalogue value from user wording.

        Values come from the loaded records, so this works for any structured
        catalogue field without maintaining a product-type keyword list.
        """
        question_terms = QueryPlanner._match_terms(question)
        matches = []
        for value in values:
            if not isinstance(value, str) or not value.strip():
                continue
            value_terms = QueryPlanner._match_terms(value)
            if value_terms and all(term in question_terms for term in value_terms):
                matches.append(value)
                continue
            if len(value_terms) == 1 and any(
                difflib.SequenceMatcher(None, value_term, question_term).ratio() >= 0.8
                for value_term in value_terms
                for question_term in question_terms
            ):
                matches.append(value)

        return matches[0] if len(matches) == 1 else None

    @staticmethod
    def _match_terms(value: Any) -> set[str]:
        terms = set(re.findall(r"[a-z0-9]+", str(value).lower()))
        return {term[:-1] if len(term) > 3 and term.endswith("s") else term for term in terms}


# ---------------------------------------------------------------------------
# Deterministic query engine
# ---------------------------------------------------------------------------

class QueryEngine:
    def __init__(self, records: List[Record]):
        self.records = records

    def execute(self, plan: Dict[str, Any]) -> List[Record]:
        records = list(self.records)

        record_ids = set(str(x) for x in plan.get("record_ids", []))
        if record_ids:
            records = [r for r in records if r.record_id in record_ids]

        for condition in plan.get("conditions", []) or []:
            records = [r for r in records if self._matches(r, condition)]

        sort = plan.get("sort")
        if sort:
            field = sort.get("field")
            reverse = sort.get("direction") == "desc"
            records.sort(
                key=lambda r: self._sort_value(r.get(field)),
                reverse=reverse,
            )

        limit = plan.get("limit")
        if isinstance(limit, int) and limit >= 0:
            records = records[:limit]

        return records

    def aggregate(self, records: List[Record], plan: Dict[str, Any]) -> Any:
        spec = plan.get("aggregate") or {}
        fn = spec.get("function")
        field = spec.get("field")

        if fn == "count":
            return len(records)

        values = [parse_number(r.get(field)) for r in records]
        values = [v for v in values if v is not None]
        if not values:
            return None

        if fn == "min":
            return min(values)
        if fn == "max":
            return max(values)
        if fn == "sum":
            return sum(values)
        if fn == "avg":
            return sum(values) / len(values)
        return None

    @staticmethod
    def _sort_value(value: Any):
        number = parse_number(value)
        if number is not None:
            return number
        return str(value or "").lower()

    def _matches(self, record: Record, condition: Dict[str, Any]) -> bool:
        field = condition.get("field")
        operator = condition.get("operator", "eq")
        expected = condition.get("value")
        actual = record.get(field)

        if actual is None:
            return False

        # Numeric comparison when possible.
        actual_num = parse_number(actual)
        expected_num = parse_number(expected)

        if actual_num is not None and expected_num is not None and operator in {"gt", "gte", "lt", "lte", "eq", "neq"}:
            left, right = actual_num, expected_num
            return {
                "gt": left > right,
                "gte": left >= right,
                "lt": left < right,
                "lte": left <= right,
                "eq": left == right,
                "neq": left != right,
            }[operator]

        left = str(actual).strip().lower()
        right = str(expected).strip().lower()

        if field == "material":
            left = canonical_material(left).lower()
            right = canonical_material(right).lower()

        if operator == "contains":
            return self._text_contains(left, right)
        if operator == "eq":
            return left == right
        if operator == "neq":
            return left != right
        if operator == "gt":
            return left > right
        if operator == "gte":
            return left >= right
        if operator == "lt":
            return left < right
        if operator == "lte":
            return left <= right
        return False

    @staticmethod
    def _text_contains(actual: str, expected: str) -> bool:
        actual_terms = QueryPlanner._match_terms(actual)
        expected_terms = QueryPlanner._match_terms(expected)
        return bool(expected_terms) and expected_terms.issubset(actual_terms)


# ---------------------------------------------------------------------------
# Semantic retrieval for unstructured questions
# ---------------------------------------------------------------------------

class SemanticRetriever:
    def __init__(self, text_records: List[Record]):
        self.records = text_records
        self.vector_store = None

        if FAISS and OpenAIEmbeddings and os.getenv("OPENAI_API_KEY") and text_records:
            try:
                docs = [
                    # Lazy import keeps structured-only deployments lightweight.
                    __import__("langchain_core.documents", fromlist=["Document"]).Document(
                        page_content=r.searchable_text(),
                        metadata={"record_id": r.record_id, "source": r.source},
                    )
                    for r in text_records
                ]
                self.vector_store = FAISS.from_documents(
                    docs,
                    OpenAIEmbeddings(model="text-embedding-3-small"),
                )
            except Exception:
                self.vector_store = None

    def search(self, query: str, k: int = 6) -> List[Record]:
        if not self.records:
            return []

        if self.vector_store:
            try:
                docs = self.vector_store.similarity_search(query, k=k)
                ids = {d.metadata.get("record_id") for d in docs}
                return [r for r in self.records if r.record_id in ids]
            except Exception:
                pass

        tokens = set(re.findall(r"[a-z0-9]+", query.lower()))
        scored = []
        for record in self.records:
            text = record.searchable_text().lower()
            score = sum(1 for token in tokens if token in text)
            if score:
                scored.append((score, record))
        scored.sort(key=lambda x: x[0], reverse=True)
        return [r for _, r in scored[:k]]


# ---------------------------------------------------------------------------
# Answer rendering
# ---------------------------------------------------------------------------

class AnswerRenderer:
    @staticmethod
    def render(records: List[Record], plan: Dict[str, Any], aggregate_value: Any = None) -> str:
        if aggregate_value is not None:
            fn = (plan.get("aggregate") or {}).get("function", "result")
            return f"{fn.title()}: {AnswerRenderer.format_value(aggregate_value)}"

        if not records:
            return "I could not find supporting information in the document catalogue."

        fields = plan.get("output_fields") or ["name"]

        lines = []
        for record in records:
            if fields == ["name"]:
                lines.append(f"- {record.get('name', record.record_id)}")
                continue

            values = []
            for field in fields:
                if field == "name":
                    continue
                value = record.get(field)
                if value in (None, ""):
                    continue
                label = field.replace("_", " ").title()
                values.append(f"{label}: {AnswerRenderer.format_value(value)}")

            if values:
                lines.append(f"- {record.get('name', record.record_id)}: " + ", ".join(values))
            else:
                lines.append(f"- {record.get('name', record.record_id)}")

        return "\n".join(lines)

    @staticmethod
    def format_value(value: Any) -> str:
        number = parse_number(value)
        if number is not None and isinstance(value, (int, float)):
            if float(number).is_integer():
                return f"₹{int(number):,}"
        return str(value)


# ---------------------------------------------------------------------------
# Main agent
# ---------------------------------------------------------------------------

class UniversalDocumentAgent:
    """
    Hybrid document agent.

    Design goals:
      1. Multiple input formats.
      2. Normalize structured information into records.
      3. Use an LLM as a query planner, not as the source of truth.
      4. Execute structured filters deterministically.
      5. Use semantic retrieval only for genuinely unstructured content.
      6. Preserve selected records for natural follow-ups.
      7. Never ask the LLM to invent catalogue values.
    """

    def __init__(self, source_path: str | Path):
        self.source_path = Path(source_path)
        self.structured_records, self.text_records = UniversalIngestor(source_path).load()

        self.engine = QueryEngine(self.structured_records)
        self.planner = QueryPlanner()
        self.semantic = SemanticRetriever(self.text_records)

        self.history: List[Dict[str, str]] = []
        self.last_record_ids: List[str] = []

    def _is_semantic_question(self, question: str, plan: Dict[str, Any]) -> bool:
        if plan.get("operation") == "semantic":
            return True
        if not self.structured_records:
            return True
        # If the planner found no structured conditions and no obvious output
        # fields, semantic retrieval may be the better route.
        return (
            not plan.get("conditions")
            and not plan.get("record_ids")
            and not plan.get("aggregate")
            and plan.get("operation") in {"semantic", "answer_from_context"}
        )

    def _followup_ids(
        self,
        question: str,
        field_values: Dict[str, List[Any]],
    ) -> List[str]:
        q = question.lower()
        if not self.last_record_ids:
            return []

        reference = bool(re.search(
            r"\b(these|those|their|them|previous|it|its)\b|\b(this|that)\s+(one|product|item)\b",
            q,
        ))
        if reference:
            return list(self.last_record_ids)

        requested_fields = {
            field for field in field_values
            if field not in {"name", "product_id"}
            and re.search(rf"\b{re.escape(field.replace('_', ' '))}\b", q)
        }
        ranking_request = any(phrase in q for phrase in [
            "highest", "lowest", "least", "most", "cheapest", "lightest", "heaviest",
        ])
        has_explicit_catalogue_value = any(
            QueryPlanner._resolve_catalogue_value(q, values) is not None
            for field, values in field_values.items()
            if field not in {"price", "weight"}
        )
        if requested_fields and not ranking_request and not has_explicit_catalogue_value:
            return list(self.last_record_ids)
        return []

    def _catalogue_schema(self) -> Dict[str, List[Any]]:
        """Return fields and their known scalar values for query planning."""
        values: Dict[str, List[Any]] = {}
        for record in self.structured_records:
            for field, value in record.fields.items():
                if value in (None, ""):
                    continue
                bucket = values.setdefault(field, [])
                if value not in bucket:
                    bucket.append(value)
        return values

    def answer(self, question: str) -> str:
        question = (question or "").strip()
        if not question:
            return "Please ask a question about the uploaded documents."

        fields = self._catalogue_schema()

        previous_ids = self._followup_ids(question, fields)
        plan = self.planner.plan(question, fields, previous_ids)

        if previous_ids:
            plan["record_ids"] = previous_ids

        if plan.get("clarify_fields"):
            response = "Which product details would you like: price, weight, material, category, or stock?"
            self.history.append({"question": question, "answer": response})
            return response

        if self.structured_records and not self._is_semantic_question(question, plan):
            records = self.engine.execute(plan)

            aggregate_value = None
            if plan.get("aggregate"):
                aggregate_value = self.engine.aggregate(records, plan)

            self.last_record_ids = [r.record_id for r in records]

            response = AnswerRenderer.render(records, plan, aggregate_value)
            self.history.append({"question": question, "answer": response})
            return response

        # Unstructured route.
        semantic_records = self.semantic.search(question)
        if not semantic_records:
            response = "I could not find supporting information in the document catalogue."
        else:
            response = self._answer_semantic(question, semantic_records)

        self.history.append({"question": question, "answer": response})
        return response

    def _answer_semantic(self, question: str, records: List[Record]) -> str:
        if not OpenAI or not os.getenv("OPENAI_API_KEY"):
            return "\n\n".join(r.searchable_text() for r in records)

        context = "\n\n".join(
            f"[Source: {r.source}]\n{r.searchable_text()}" for r in records
        )
        prompt = f"""
Answer the user's question using ONLY the supplied source context.

Rules:
- Do not invent facts.
- Do not fill missing fields from general knowledge.
- If the context does not support the answer, say that the supporting information
  was not found in the document catalogue.
- Keep the answer concise.
- Do not claim that you checked documents that are not in the context.

Source context:
{context}

User question:
{question}
"""
        try:
            client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
            response = client.responses.create(
                model=os.getenv("OPENAI_ANSWER_MODEL", "gpt-4o-mini"),
                input=[{"role": "user", "content": prompt}],
            )
            return response.output_text.strip()
        except Exception:
            return "\n\n".join(r.searchable_text() for r in records)


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Universal hybrid document agent")
    parser.add_argument("source", help="PDF/file/folder containing source documents")
    args = parser.parse_args()

    agent = UniversalDocumentAgent(args.source)
    print("Universal Document Agent")
    print("Type 'exit' to quit.")
    while True:
        question = input("\nYou: ").strip()
        if question.lower() in {"exit", "quit"}:
            break
        print("Agent:", agent.answer(question))
